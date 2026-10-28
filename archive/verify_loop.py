"""回归测试：逐条覆盖群评审提出的问题。

跑法：python verify_loop.py
任何一个断言失败都说明评审指出的问题没修好。
"""

from __future__ import annotations

import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from engine import (
    STATUS_CONFIRMED_SAME,
    STATUS_CONFLICT,
    STATUS_NEED_HUMAN,
    ContractDataError,
    ContractRate,
    DiffStatus,
    LineItem,
    RateStatus,
    build_rate_index,
    item_key,
    load_contract_rates,
    load_line_items,
    load_plan,
    load_plan_with_display,
    reconcile,
    resolve_rate,
)
from rulebook import (
    DIFF_BEFORE_CONTRACT,
    DIFF_DATE_RANGE,
    DIFF_DATE_UNPARSABLE,
    DIFF_EXTRA_ROW,
    DIFF_MISSING_ROW,
    DIFF_NO_CONTRACT,
    DIFF_UNIT_PRICE,
    VERDICT_ACCEPT_SUGGESTED,
    VERDICT_KEEP_ORIGINAL,
    Rulebook,
)

BASE = Path(__file__).resolve().parent
SAMPLE = BASE / "sample"
CUSTOMER = "介護グループA"

passes = 0
failures: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    global passes
    if cond:
        passes += 1
        print(f"  PASS  {name}")
    else:
        failures.append(name)
        print(f"  FAIL  {name}  {detail}")


def by_type(diffs, t):
    return [d for d in diffs if d.diff_type == t]


# ---------------------------------------------------------------- 问题 ①
print("\n【问题①】合同单价表重复键不得静默覆盖")
dup_path = BASE / "tmp_dup_contract.csv"
dup_path.write_text(
    "利用者名,品目,契約単価\n"
    "鈴木 次郎,訪問介護（生活援助2),1250\n"
    "鈴木 次郎,訪問介護（生活援助2),1320\n",
    encoding="utf-8",
)
try:
    load_contract_rates(dup_path)
    check("无生效日的重复键应报错", False, "竟然没报错，被静默覆盖了")
except ContractDataError as e:
    check("无生效日的重复键应报错", True)
    print(f"        报错信息: {e}")

same_eff = BASE / "tmp_same_eff.csv"
same_eff.write_text(
    "利用者名,品目,契約単価,適用開始日\n"
    "鈴木 次郎,訪問介護（生活援助2),1250,2026-04-01\n"
    "鈴木 次郎,訪問介護（生活援助2),1320,2026-04-01\n",
    encoding="utf-8",
)
try:
    load_contract_rates(same_eff)
    check("相同生效日的重复键应报错", False, "竟然没报错")
except ContractDataError:
    check("相同生效日的重复键应报错", True)

segmented = load_contract_rates(SAMPLE / "契約単価表_分段.csv")
seg_idx = build_rate_index(segmented)
r_oct = resolve_rate(seg_idx, "鈴木 次郎", "訪問介護（生活援助2)", "2026-10-20")
check(
    "合法的分段单价（不同生效日）应能加载",
    r_oct.status is RateStatus.FOUND and r_oct.price == 1250,
    f"实际 {r_oct}",
)
r_nov = resolve_rate(seg_idx, "鈴木 次郎", "訪問介護（生活援助2)", "2026-11-05")
check("11 月起应取新单价 1,320", r_nov.status is RateStatus.FOUND and r_nov.price == 1320, f"实际 {r_nov}")

# 「早于所有分段」必须能区分出来，不能被当成合同未约定或数据质量问题
r_before = resolve_rate(seg_idx, "鈴木 次郎", "訪問介護（生活援助2)", "2026-01-05")
check(
    "服务日早于合同所有分段应单独识别",
    r_before.status is RateStatus.BEFORE_ALL,
    f"实际 {r_before.status}",
)
check("应带出合同最早生效日", r_before.earliest_from == "2026-04-01", f"实际 {r_before.earliest_from}")
r_none = resolve_rate(seg_idx, "不存在的 人", "訪問介護（身体介護1)", "2026-10-05")
check("合同里没这个品目应报 NO_CONTRACT", r_none.status is RateStatus.NO_CONTRACT)

# 排序后必须严格递增（Chief/CodeReviewer 定的 B1 口径）
segs_after_sort = seg_idx[item_key("鈴木 次郎", "訪問介護（生活援助2)")]
effs_sorted = [s.effective_from for s in segs_after_sort]
check(
    "分段排序后生效日严格递增",
    effs_sorted == sorted(set(effs_sorted)) and len(effs_sorted) == len(set(effs_sorted)),
    f"实际 {effs_sorted}",
)

# 乱序输入也要能正确排序（客户 Excel 里常见顺序混乱）
shuffled = BASE / "tmp_shuffled.csv"
shuffled.write_text(
    "利用者名,品目,契約単価,適用開始日\n"
    "鈴木 次郎,訪問介護（生活援助2),1320,2026-10-01\n"
    "鈴木 次郎,訪問介護（生活援助2),1250,2026-04-01\n",
    encoding="utf-8",
)
sh_idx = build_rate_index(load_contract_rates(shuffled))
sh_segs = sh_idx[item_key("鈴木 次郎", "訪問介護（生活援助2)")]
check(
    "乱序输入应被正确排序",
    [s.effective_from for s in sh_segs] == ["2026-04-01", "2026-10-01"],
    f"实际 {[s.effective_from for s in sh_segs]}",
)
check(
    "乱序输入取 10 月价应是新价",
    resolve_rate(sh_idx, "鈴木 次郎", "訪問介護（生活援助2)", "2026-10-15").price == 1320,
)
shuffled.unlink()

# 生效日不是合法日期时必须报错，否则排序不可信
bad_eff = BASE / "tmp_bad_eff.csv"
bad_eff.write_text(
    "利用者名,品目,契約単価,適用開始日\n"
    "佐藤 花子,訪問介護（身体介護1),1830,2026/04/01\n",
    encoding="utf-8",
)
try:
    load_contract_rates(bad_eff)
    check("非法生效日应报错", False, "竟然没报错")
except ContractDataError:
    check("非法生效日应报错", True)
bad_eff.unlink()
for p in (dup_path, same_eff):
    p.unlink()

# ---------------------------------------------------------------- 问题 ②
print("\n【问题②】缺失行不得用合同全量判定")
full_contract = load_contract_rates(SAMPLE / "契約単価表.csv")  # 含山田，全量 10 行
full_idx = build_rate_index(full_contract)
items_oct = load_line_items(SAMPLE / "利用実績_2026-10.csv")

diffs_no_plan, warns = reconcile(items_oct, full_idx, "2026-10", plan=None)
check(
    "无计划表时不产出缺失行",
    len(by_type(diffs_no_plan, DIFF_MISSING_ROW)) == 0,
    f"却产出了 {len(by_type(diffs_no_plan, DIFF_MISSING_ROW))} 条",
)
check("无计划表时必须给出警告", any("缺失行" in w for w in warns), f"warnings={warns}")

plan = load_plan(SAMPLE / "計画明細_2026-10.csv")
diffs_with_plan, warns2 = reconcile(items_oct, full_idx, "2026-10", plan=plan)
check(
    "有计划表时不出现虚假缺失行（山田不在计划内）",
    len(by_type(diffs_with_plan, DIFF_MISSING_ROW)) == 0,
    f"却产出了 {len(by_type(diffs_with_plan, DIFF_MISSING_ROW))} 条",
)

# ---------------------------------------------------------------- 问题 ③
print("\n【问题③】指纹不得泛化，三级判定")
check(
    "单价差异指纹不含客户/品目以外的金额维度则不合格",
    all(
        "1,300" in d.fingerprint or "1300" in d.fingerprint
        for d in by_type(diffs_with_plan, DIFF_UNIT_PRICE)
    ),
    "指纹里没带原单价，会泛化",
)

book = Rulebook(customer=CUSTOMER)
# 鈴木 特例单价 1,300 被客户认可 -> 保留原值
unit_d = by_type(diffs_with_plan, DIFF_UNIT_PRICE)[0]
book.confirm(unit_d.diff_type, unit_d.fingerprint, VERDICT_KEEP_ORIGINAL, "2026-10",
             note="特例単価、客户认可")

# 次月：同样的 1,300 应命中「已确认·建议不变」
diffs_same, _ = reconcile(items_oct, full_idx, "2026-10", rulebook=book, plan=plan)
d_same = by_type(diffs_same, DIFF_UNIT_PRICE)[0]
check(
    "同数值重复出现 -> 已确认过·建议不变",
    d_same.status == STATUS_CONFIRMED_SAME,
    f"实际 {d_same.status}",
)

# 关键场景：下月铃木真的录错成 1,500，不能被规则掩盖
bad_items = [
    LineItem(
        row_no=6,
        customer_name="鈴木 次郎",
        item_name="訪問介護（生活援助2)",
        quantity=5,
        unit_price=1500,
        amount=7500,
        service_date="2026-10-20",
    )
]
diffs_err, _ = reconcile(bad_items, full_idx, "2026-10", rulebook=book, plan=plan)
d_err = by_type(diffs_err, DIFF_UNIT_PRICE)
check("真实的 1,500 错录不应命中旧规则", len(d_err) == 1)
check(
    "真实的 1,500 错录应标为需确认",
    d_err and d_err[0].status == STATUS_NEED_HUMAN,
    f"实际 {d_err[0].status if d_err else '无差异'}",
)

# 上次「采纳建议值」的项再次出现 -> 上次已改、本次再现（RECURRING），不是「与上次判定不符」
book2 = Rulebook(customer=CUSTOMER)
book2.confirm(unit_d.diff_type, unit_d.fingerprint, VERDICT_ACCEPT_SUGGESTED, "2026-10")
diffs_c, _ = reconcile(items_oct, full_idx, "2026-10", rulebook=book2, plan=plan)
d_c = by_type(diffs_c, DIFF_UNIT_PRICE)[0]
check(
    "上次采纳建议值、本次同值再现 -> 上次已改，本次再现",
    d_c.status is DiffStatus.RECURRING,
    f"实际 {d_c.status}",
)
check(
    "RECURRING 仍计入需人看",
    d_c.needs_human,
    "重复出错的项不能被静默放过",
)

# ---------------------------------------------------------------- 小问题：日期
print("\n【小问题】日期无法解析必须报出，不能放过")
bad_date_items = [
    LineItem(
        row_no=1,
        customer_name="佐藤 花子",
        item_name="訪問介護（身体介護1)",
        quantity=1,
        unit_price=1830,
        amount=1830,
        service_date="令和8年10月5日",
    )
]
diffs_bd, _ = reconcile(bad_date_items, full_idx, "2026-10", plan=plan)
check(
    "非 ISO 日期应报「日期无法解析」",
    len(by_type(diffs_bd, DIFF_DATE_UNPARSABLE)) == 1,
    f"实际 {len(by_type(diffs_bd, DIFF_DATE_UNPARSABLE))} 条",
)

# ---------------------------------------------------------------- 清单 v1 第 7 类
print("\n【清单v1·第7类】合同未约定不得静默放过")
no_contract_items = [
    LineItem(
        row_no=1,
        customer_name="佐藤 花子",
        item_name="緊急訪問介護（臨時加算）",  # 合同表里没有这个品目
        quantity=1,
        unit_price=5000,
        amount=5000,
        service_date="2026-10-05",
    )
]
diffs_nc, _ = reconcile(no_contract_items, full_idx, "2026-10", plan=plan)
nc = by_type(diffs_nc, DIFF_NO_CONTRACT)
check("合同查不到的行必须报「合同未约定」", len(nc) == 1, f"实际 {len(nc)} 条")
check(
    "合同未约定置信度应为 0.6",
    nc and abs(nc[0].confidence - 0.6) < 1e-9,
    f"实际 {nc[0].confidence if nc else '无'}",
)
check(
    "合同未约定不能同时报成单价不一致",
    len(by_type(diffs_nc, DIFF_UNIT_PRICE)) == 0,
)

# 第 7 类：首次必须人工确认，但确认后应当能被记住（指纹不含金额就是为了这个）
book_nc = Rulebook(customer=CUSTOMER)
book_nc.confirm(nc[0].diff_type, nc[0].fingerprint, VERDICT_KEEP_ORIGINAL, "2026-10")

# 下月同一合同外服务，金额变了，仍然必须命中同一条规则（不能被当成新问题再问一遍）
nc_next_amount = [
    LineItem(
        row_no=1,
        customer_name="佐藤 花子",
        item_name="緊急訪問介護（臨時加算）",
        quantity=2,
        unit_price=5000,
        amount=10000,  # 金额变了，指纹不该变
        service_date="2026-11-05",
    )
]
diffs_nc2, _ = reconcile(nc_next_amount, full_idx, "2026-11", rulebook=book_nc, plan=set())
nc2 = by_type(diffs_nc2, DIFF_NO_CONTRACT)
check(
    "合同未约定确认过之后，下月金额变化仍能命中（规则簿学得会）",
    nc2 and nc2[0].status is DiffStatus.CONFIRMED_SAME,
    f"实际 {nc2[0].status if nc2 else '无'}",
)
check(
    "但首次出现时必须是需确认",
    nc[0].status is DiffStatus.NEW,
)

# 日期无法解析则不同：那是数据格式问题，格式没修就不该被记住
book_du = Rulebook(customer=CUSTOMER)
du_fp = by_type(diffs_bd, DIFF_DATE_UNPARSABLE)[0].fingerprint
book_du.confirm(DIFF_DATE_UNPARSABLE, du_fp, VERDICT_KEEP_ORIGINAL, "2026-10")
diffs_du2, _ = reconcile(
    bad_date_items,
    full_idx,
    "2026-10",
    rulebook=book_du,
    plan={item_key("佐藤 花子", "訪問介護（身体介護1)")},
)
check(
    "日期无法解析永远需确认，不被规则簿记住",
    by_type(diffs_du2, DIFF_DATE_UNPARSABLE)[0].status is DiffStatus.NEW,
    f"实际 {by_type(diffs_du2, DIFF_DATE_UNPARSABLE)[0].status}",
)

# 第 7 类指纹不得带金额：金额每月都会变，带进去规则永远学不会
nc_fp = nc[0].fingerprint
check(
    "第 7 类指纹不含金额",
    "5,000" not in nc_fp and "5000" not in nc_fp,
    f"指纹是 {nc_fp}",
)
# 同一合同外服务，下月金额变了，指纹必须不变
next_month_nc = [
    LineItem(
        row_no=1,
        customer_name="佐藤 花子",
        item_name="緊急訪問介護（臨時加算）",
        quantity=2,
        unit_price=5000,
        amount=10000,  # 金额变了
        service_date="2026-11-05",
    )
]
diffs_nc3, _ = reconcile(next_month_nc, full_idx, "2026-11", plan=set())
nc3 = by_type(diffs_nc3, DIFF_NO_CONTRACT)
check("下月金额变化，第 7 类指纹应保持不变", nc3 and nc3[0].fingerprint == nc_fp,
      f"实际 {nc3[0].fingerprint if nc3 else '无'}，原为 {nc_fp}")

# ---------------------------------------------------------------- 清单 v1 第 4 类
print("\n【清单v1·第4类】多余行必须带具体理由")
extra_items = load_line_items(SAMPLE / "利用実績_2026-10.csv") + [
    LineItem(
        row_no=99,
        customer_name="佐藤 花子",
        item_name="通院介助",  # 不在计划表，但合同里有
        quantity=1,
        unit_price=2000,
        amount=2000,
        service_date="2026-10-08",
    )
]
# 让「通院介助」出现在合同里，才能测「合同有、计划无」这种真正的多余行
rates_with_extra = load_contract_rates(SAMPLE / "契約単価表_分段.csv") + [
    ContractRate(customer_name="佐藤 花子", item_name="通院介助", unit_price=2000)
]
idx_extra = build_rate_index(rates_with_extra)
diffs_ex, _ = reconcile(extra_items, idx_extra, "2026-10", plan=plan)
ex = by_type(diffs_ex, DIFF_EXTRA_ROW)
check("计划外行必须报「多余行」", len(ex) == 1, f"实际 {len(ex)} 条")
check(
    "多余行必须给出具体理由而不只是类型",
    ex and "不在本月计划内" in ex[0].suggested,
    f"实际理由 {ex[0].suggested if ex else '无'}",
)

# 合同里也没有的行：报「合同未约定」即可，不应再叠加一条多余行
diffs_dup, _ = reconcile(no_contract_items, full_idx, "2026-10", plan=plan)
check(
    "合同未约定的行不重复报成多余行",
    len(by_type(diffs_dup, DIFF_EXTRA_ROW)) == 0,
    f"却报了 {len(by_type(diffs_dup, DIFF_EXTRA_ROW))} 条多余行",
)

# ---------------------------------------------------------------- 清单 v1 第 8 类
print("\n【清单v1·第8类】日期无法解析不计入「N 处不一致」")
from report import write_report  # noqa: E402

report_path = BASE / "out" / "_tmp_split.xlsx"
# 用只含这一行的计划表，隔离出「日期无法解析」这一类的影响
# 注意：plan 现在一律是归一化键
single_plan = {item_key("佐藤 花子", "訪問介護（身体介護1)")}
diffs_bd_iso, _ = reconcile(bad_date_items, full_idx, "2026-10", plan=single_plan)
write_report(diffs_bd_iso, report_path, CUSTOMER, "2026-10")
from openpyxl import load_workbook  # noqa: E402

wsv = load_workbook(report_path).active
summary = wsv["A2"].value
check(
    "总数不应包含日期无法解析的行",
    summary is not None and "本月 0 处不一致" in summary,
    f"实际摘要 {summary}",
)
all_text = " ".join(
    str(c.value) for r in wsv.iter_rows() for c in r if c.value is not None
)
check("数据质量问题应单列段落", "数据质量问题" in all_text)
report_path.unlink()

# ---------------------------------------------------------------- 品目名归一化
print("\n【归一化】写法差异不得被误报成缺失行/多余行")
from engine import item_key, norm_item, norm_name  # noqa: E402  (item_key 已在顶部导入)

check("全角数字应归一为半角", norm_item("訪問介護（身体介護１）") == "訪問介護(身体介護1)", norm_item("訪問介護（身体介護１）"))
check("半角右括号应归一", norm_item("訪問介護（身体介護1)") == "訪問介護(身体介護1)", norm_item("訪問介護（身体介護1)"))
check("中间空格应去除", norm_item("訪問介護 （身体介護1）") == "訪問介護(身体介護1)", norm_item("訪問介護 （身体介護1）"))
check("全角空格应去除", norm_name("鈴木　次郎") == "鈴木次郎", norm_name("鈴木　次郎"))
check("敬称「様」应去除", norm_name("鈴木 次郎様") == "鈴木次郎", norm_name("鈴木 次郎様"))
check(
    "不同品目不得归一化成一回事（身体介護1 vs 2）",
    norm_item("訪問介護（身体介護1）") != norm_item("訪問介護（身体介護2）"),
)
check(
    "「介護」和「介护」确实是不同字，不强行合并",
    norm_item("介護") != norm_item("介护"),
    "这两者是不同字符，归一化不该做语义翻译",
)

# 真实场景：请求书用半角括号+全角数字，合同用全角——应当匹配上，不报差异
write_items = [
    LineItem(
        row_no=1,
        customer_name="佐藤 花子",
        item_name="訪問介護（身体介護１)",  # 全角数字 + 半角括号
        quantity=12,
        unit_price=1830,
        amount=21960,
        service_date="2026-10-05",
    )
]
diffs_norm, _ = reconcile(write_items, full_idx, "2026-10", plan=None)
check("写法差异不应报成「合同未约定」", len(by_type(diffs_norm, DIFF_NO_CONTRACT)) == 0,
      f"却报了 {len(by_type(diffs_norm, DIFF_NO_CONTRACT))} 条")
check("写法差异不应报成单价不一致", len(by_type(diffs_norm, DIFF_UNIT_PRICE)) == 0)

# 缺失行/多余行也要走归一化
norm_plan = {item_key("佐藤 花子", "訪問介護（身体介護1)")}
diffs_norm2, _ = reconcile(write_items, full_idx, "2026-10", plan=norm_plan)
check("写法差异不应报成缺失行", len(by_type(diffs_norm2, DIFF_MISSING_ROW)) == 0,
      f"却报了 {len(by_type(diffs_norm2, DIFF_MISSING_ROW))} 条缺失行")
check("写法差异不应报成多余行", len(by_type(diffs_norm2, DIFF_EXTRA_ROW)) == 0,
      f"却报了 {len(by_type(diffs_norm2, DIFF_EXTRA_ROW))} 条多余行")

# 真实场景回归：一份写法故意不统一的数据，不应被刷屏
mixed_items = load_line_items(SAMPLE / "利用実績_写法混乱.csv")
mixed_plan, mixed_disp = load_plan_with_display(SAMPLE / "計画明細_2026-10.csv")
diffs_mixed, _ = reconcile(
    mixed_items, full_idx, "2026-10", plan=mixed_plan, plan_display=mixed_disp
)
check(
    "写法混乱的数据不应刷出假缺失行",
    len(by_type(diffs_mixed, DIFF_MISSING_ROW)) == 0,
    f"却报了 {len(by_type(diffs_mixed, DIFF_MISSING_ROW))} 条",
)
check(
    "写法混乱的数据不应刷出假多余行",
    len(by_type(diffs_mixed, DIFF_EXTRA_ROW)) == 0,
    f"却报了 {len(by_type(diffs_mixed, DIFF_EXTRA_ROW))} 条",
)
check(
    "写法混乱的数据仍应报出真实的那一处差异",
    len(diffs_mixed) == 1 and diffs_mixed[0].diff_type == DIFF_UNIT_PRICE,
    f"实际 {[(d.diff_type) for d in diffs_mixed]}",
)

# ---------------------------------------------------------------- 日期早于合同生效
print("\n【新】日期早于合同所有分段必须报出，不能被跳过")
old_period_items = [
    LineItem(
        row_no=1,
        customer_name="鈴木 次郎",
        item_name="訪問介護（生活援助2)",
        quantity=5,
        unit_price=1250,
        amount=6250,
        service_date="2026-01-15",  # 合同最早 2026-04-01 生效
    )
]
diffs_before, _ = reconcile(old_period_items, seg_idx, "2026-01", plan=set())
bc = by_type(diffs_before, DIFF_BEFORE_CONTRACT)
check("日期早于合同起点应报「日期早于合同生效」", len(bc) == 1, f"实际 {len(bc)} 条")
check(
    "应指出合同最早生效日",
    bc and "2026-04-01" in bc[0].suggested,
    f"实际 {bc[0].suggested if bc else '无'}",
)
check(
    "子原因要在文案里标出「合同未覆盖该日期」，与普通合同外服务区分",
    bc and "合同未覆盖该日期" in bc[0].suggested,
    f"实际 {bc[0].suggested if bc else '无'}",
)
check(
    "早于合同生效的指纹与普通合同未约定前缀不同（语义相反，不能共用一个）",
    bc and bc[0].fingerprint.startswith("before-contract"),
    f"实际指纹 {bc[0].fingerprint if bc else '无'}",
)
check(
    "早于合同生效的指纹不含金额（能被规则簿索引）",
    bc and len(bc[0].fingerprint.split("|")) == 3,
    f"实际指纹 {bc[0].fingerprint if bc else '无'}",
)
# 内部全集必须覆盖它，否则按 ALL_DIFF_TYPES 遍历的代码会静默漏掉这一类
from rulebook import ALL_DIFF_TYPES, ALL_INTERNAL_DIFF_TYPES  # noqa: E402

check(
    "代码内部全集包含「日期早于合同生效」",
    DIFF_BEFORE_CONTRACT in ALL_INTERNAL_DIFF_TYPES,
    f"实际 {ALL_INTERNAL_DIFF_TYPES}",
)
check(
    "对外 8 类不含它（客户只看到 8 类）",
    DIFF_BEFORE_CONTRACT not in ALL_DIFF_TYPES,
)
check(
    "不应重复报「日期不在期间内」（说的是同一件事）",
    len(by_type(diffs_before, DIFF_DATE_RANGE)) == 0,
    f"却报了 {len(by_type(diffs_before, DIFF_DATE_RANGE))} 条",
)

# ---------------------------------------------------------------- 类型口径
print("\n【类型口径】保持 8 类，「早于合同生效」并入第 7 类用子原因区分")
from rulebook import ALL_DIFF_TYPES  # noqa: E402

check("对外口径仍是 8 类", len(ALL_DIFF_TYPES) == 8, f"实际 {len(ALL_DIFF_TYPES)} 类")

# ---------------------------------------------------------------- 三态枚举
print("\n【三态枚举】DiffStatus 必须是真的枚举，不是布尔打补丁")
from engine import DiffStatus  # noqa: E402

check("状态必须是枚举成员", isinstance(d_err[0].status, DiffStatus), f"实际类型 {type(d_err[0].status)}")
check("NEW 计入需人看", DiffStatus.NEW != DiffStatus.CONFIRMED_SAME)
check("枚举共四态", len(list(DiffStatus)) == 4, f"实际 {len(list(DiffStatus))} 态")
check(
    "RECURRING 的文案是「上次已改，本次再现」，不是「与上次判定不符」",
    "上次已改" in DiffStatus.RECURRING.value and "不符" not in DiffStatus.RECURRING.value,
    f"实际文案 {DiffStatus.RECURRING.value}",
)

# ---------------------------------------------------------------- 汇总
print("\n" + "=" * 60)
print(f"通过 {passes} 项，失败 {len(failures)} 项")
if failures:
    for f in failures:
        print(f"  ✗ {f}")
    raise SystemExit(1)
print("全部通过：CodeReviewer 指出的问题已修复。")
